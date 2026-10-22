"""Verify a rendered manuscript without modifying source templates or results."""

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader


parser = argparse.ArgumentParser()
parser.add_argument("--pdf", type=Path, required=True, help="Rendered manuscript PDF")
parser.add_argument("--report", type=Path, help="Optional path for a verification report")
args = parser.parse_args()

root = Path(__file__).resolve().parents[1]
manuscript = Document(root / "FA_MAS_TSC_Submission_Ready_Manuscript.docx")
template = Document(next((root / "official_template").glob("*.docx")))
styles = ["Normal", "Title", "Heading 1", "Heading 2", "Text", "Abstract"]
for name in styles:
    assert manuscript.styles[name].element.xml == template.styles[name].element.xml, name

assert args.pdf.is_file(), args.pdf
pages = PdfReader(args.pdf).pages
assert len(pages) <= 14, len(pages)
assert len(pages) == 13, len(pages)
text = " ".join(" ".join((page.extract_text() or "").split()) for page in pages)
for phrase in [
    "OpenAI GPT-6 Luna",
    "64 authored local-service trajectories",
    "3.125-point",
    "one frozen named-configuration live gate ablation",
    "cannot establish production readiness",
]:
    assert phrase in text, phrase

abstract = next(paragraph.text for paragraph in manuscript.paragraphs if paragraph.text.startswith("Abstract"))
abstract_words = len(abstract.split())
assert abstract_words <= 250, abstract_words

result = {
    "pages": len(pages),
    "within_14_page_ceiling": True,
    "abstract_words": abstract_words,
    "template_styles_identical": styles,
    "pdf_sha256": hashlib.sha256(args.pdf.read_bytes()).hexdigest(),
    "live_outcomes_claimed": True,
    "author_declarations_confirmed": False,
}
if args.report:
    args.report.write_text(json.dumps(result, indent=2) + "\n")
print(f"Verified {len(pages)} pages, {abstract_words} abstract words, and six template styles.")
